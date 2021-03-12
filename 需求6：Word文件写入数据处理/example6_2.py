import openpyxl
import os
import docx
from docx.document import Document

from openpyxl.worksheet.worksheet import Worksheet

# 模板路径
docTemplatePath = "./resources/ex6_2/八周年邀请函模版.docx"
# 花名册路径
nameListPath = "./resources/ex6_2/员工花名册.xlsx"

# 邀请函所在目录
inviteListPath = "./resources/ex6_2/邀请函"
# 没有要创建
if not os.path.exists(path=inviteListPath):
    os.mkdir(path=inviteListPath)

nameList = []
# 通过花名册导出姓名
nameSheet = openpyxl.load_workbook(filename=nameListPath)["汇总"]  # type: Worksheet
for rowData in nameSheet.iter_rows(min_row=2):
    name = rowData[0].value
    nameList.append(name)

# 遍历姓名，通过模板创建邀请函
for name in nameList:
    # 每趟遍历后必须新建，不然内容会叠加
    tempDoc = docx.Document(docx=docTemplatePath)  # type: Document
    # 组装完整文件名
    invitePath = os.path.join(inviteListPath, f"八周年邀请函_致{name}.docx")
    # 指定样式处插入文本
    tempDoc.paragraphs[3].runs[0].add_text(text=name)
    # 另存一份，不影响原来路径
    tempDoc.save(path_or_stream=invitePath)