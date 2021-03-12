import openpyxl
import docx
import os

from docx.document import Document
from openpyxl import Workbook
from openpyxl.cell import Cell

os.chdir(path="./resources/ex6_1")

docPath = "年报.docx"
excelPath = "年报.xlsx"

# 读表文件
reportBook = openpyxl.load_workbook(filename=excelPath, data_only=True)  # type: Workbook

# 读文档
reportDoc = docx.Document(docx=docPath)  # type: Document

# 任务：excel表复制到word表中去
for sheetIndex, curSheet in enumerate(reportBook.worksheets):
    # 遍历excel表的行
    for rowIndex, rowData in enumerate(curSheet.rows):
        # 遍历excel当前行的所有列
        for colIndex, cellData in enumerate(rowData):
            # word表中填入数据
            content = cellData.value
            # 整型还要分情况做一下处理
            if content is None:
                continue
            elif type(content) == int:
                # 负数先转整数，上括号和千位符
                if content < 0:
                    content = f"({abs(content):,})"
                else:
                    content = f"{content:,}"

            # 提取word文档中对应的表
            curTable = reportDoc.tables[sheetIndex]
            # 添加相应值
            curTable.cell(row_idx=rowIndex, col_idx=colIndex).text = content

# 保存word文档
reportDoc.save(path_or_stream="年报_改.docx")