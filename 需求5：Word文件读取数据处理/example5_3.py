import openpyxl
import docx
from docx.document import Document
from docx.text.run import Run
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

# 单词表

wordListPath = "./resources/ex5_3/数学英语词汇.docx"
# 默写词汇
memorizeListPath = "./resources/ex5_3/重点单词.xlsx"

# 开簿
phraseBook = openpyxl.load_workbook(filename=memorizeListPath)  # type: Workbook
# 开表
phraseSheet = phraseBook["数学"]  # type: Worksheet

# 开文档
phraseDoc = docx.Document(docx=wordListPath)  # type: Document

# 遍历文档，提取正确答案
correctAnswers = []
for paragraph in phraseDoc.paragraphs:
    for run in paragraph.runs:
        if run.bold or run.underline:
            correctAnswers.append(paragraph.runs[0].text)

# 匹配表格，对答案
for rowIndex, rowData in enumerate(phraseSheet.iter_rows(max_col=3)):
    if correctAnswers[rowIndex] != rowData[1].value:
        # 将正解加入C列
        rowData[2].value = correctAnswers[rowIndex]

# 保存表格
phraseBook.save(filename=memorizeListPath)