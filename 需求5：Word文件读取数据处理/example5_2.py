import docx
from docx.document import Document
from docx.text.run import Run
import openpyxl
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

phrasePath = "./resources/ex5_2/数学英语词汇.docx"
# 目标路径
xlsxTargetPath = "./resources/ex5_2/重点单词.xlsx"

# 开文档
phrase_doc = docx.Document(docx=phrasePath)  # type: Document
# 开表格
phrase_xlsx = openpyxl.Workbook()  # type: Workbook

# 取工作表单
mathSheet = phrase_xlsx["Sheet"]  # type: Worksheet
mathSheet.title = "数学"

# 遍历文档
for paragraph in phrase_doc.paragraphs:
    for run in paragraph.runs:
        if run.bold or run.underline:
            mathSheet.append([run.text])

# 保存表格
phrase_xlsx.save(filename=xlsxTargetPath)